"""Confirmed meeting protocols rendered locally as PDF and DOCX."""

import io
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from fastapi import HTTPException
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, TableStyle
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.models import ActionItem, Meeting, MeetingSegment, MeetingSpeaker, Protocol


@dataclass
class ProtocolDocument:
    title: str
    meeting_date: date
    speakers: list[tuple[str, str]]
    summary: str
    decisions: list[str]
    action_items: list[dict]
    transcript: list[dict]


def _deadline(item: dict) -> str:
    spoken, resolved = item.get("deadline_text") or "", str(item.get("deadline_date") or "")
    return f"{spoken} ({resolved})" if spoken and resolved else spoken or resolved or "срок не указан"


def _cells(item: dict) -> list[str]:
    return [
        str(item.get("text") or ""),
        str(item.get("owner_name") or "не назначен"),
        _deadline(item),
        str(item.get("urgency") or "средний"),
    ]


_HEADINGS = ["Поручение", "Ответственный", "Срок", "Приоритет"]


def to_pdf(doc: ProtocolDocument) -> bytes:
    fonts = Path(__file__).with_name("fonts")
    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(fonts / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(fonts / "DejaVuSans-Bold.ttf")))
    body = ParagraphStyle("body", fontName="DejaVu", fontSize=9.5, leading=14, spaceAfter=6)
    h1 = ParagraphStyle("title", parent=body, fontName="DejaVu-Bold", fontSize=18, leading=23, spaceAfter=12)
    h2 = ParagraphStyle(
        "section", parent=body, fontName="DejaVu-Bold", fontSize=12, leading=16, spaceBefore=10, keepWithNext=True
    )

    def paragraph(text, style=body):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    flow = [paragraph(doc.title, h1), paragraph(doc.meeting_date.strftime("%d.%m.%Y")), paragraph("Участники", h2)]
    flow.extend(paragraph(f"{speaker}: {name}") for speaker, name in doc.speakers)
    flow.extend([paragraph("Саммари", h2), paragraph(doc.summary), paragraph("Решения", h2)])
    flow.extend(paragraph(f"{i}. {text}") for i, text in enumerate(doc.decisions, 1))
    flow.append(paragraph("Поручения", h2))
    if doc.action_items:
        rows = [[paragraph(label) for label in _HEADINGS]]
        rows.extend([paragraph(value) for value in _cells(item)] for item in doc.action_items)
        table = LongTable(rows, colWidths=[72 * mm, 38 * mm, 41 * mm, 23 * mm], repeatRows=1, splitInRow=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef4")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#a0acb8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        flow.append(table)
    else:
        flow.append(paragraph("Поручения не зафиксированы."))
    flow.append(paragraph("Стенограмма", h2))
    for segment in doc.transcript:
        flow.append(paragraph(f"[{segment['t']}] {segment['speaker']}: {segment['text']}"))
    output = io.BytesIO()
    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=doc.title,
        author="Protokol",
    )

    def page_number(canvas, document):
        canvas.saveState()
        canvas.setFont("DejaVu", 8)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(document.page))
        canvas.restoreState()

    pdf.build(flow, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()


def to_docx(doc: ProtocolDocument) -> bytes:
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(doc.title, 0)
    document.add_paragraph(doc.meeting_date.strftime("%d.%m.%Y"))
    document.add_heading("Участники", 1)
    for speaker, name in doc.speakers:
        document.add_paragraph(f"{speaker}: {name}")
    document.add_heading("Саммари", 1)
    document.add_paragraph(doc.summary)
    document.add_heading("Решения", 1)
    for decision in doc.decisions:
        document.add_paragraph(decision, style="List Bullet")
    document.add_heading("Поручения", 1)
    table = document.add_table(rows=1, cols=4, style="Table Grid")
    for cell, label in zip(table.rows[0].cells, _HEADINGS, strict=True):
        cell.text = label
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for item in doc.action_items:
        for cell, value in zip(table.add_row().cells, _cells(item), strict=True):
            cell.text = value
    document.add_heading("Стенограмма", 1)
    for segment in doc.transcript:
        document.add_paragraph(f"[{segment['t']}] {segment['speaker']}: {segment['text']}")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


async def build_document(session: AsyncSession, meeting_id: str) -> ProtocolDocument:
    meeting = await session.get(Meeting, meeting_id)
    if meeting is None:
        raise HTTPException(404, "Совещание не найдено")
    protocol = await session.scalar(
        select(Protocol)
        .where(Protocol.meeting_id == meeting_id, Protocol.confirmed_at.is_not(None))
        .order_by(Protocol.confirmed_at.desc(), Protocol.id.desc())
        .limit(1)
    )
    if protocol is None:
        raise HTTPException(404, "Подтверждённый протокол не найден")
    speakers = list(
        await session.scalars(
            select(MeetingSpeaker).where(MeetingSpeaker.meeting_id == meeting_id).order_by(MeetingSpeaker.speaker_id)
        )
    )
    names = {speaker.speaker_id: speaker.display_name for speaker in speakers}
    segments = list(
        await session.scalars(
            select(MeetingSegment).where(MeetingSegment.meeting_id == meeting_id).order_by(MeetingSegment.idx)
        )
    )
    items = list(
        await session.scalars(
            select(ActionItem).where(ActionItem.protocol_id == protocol.id).order_by(ActionItem.action_key)
        )
    )
    return ProtocolDocument(
        title=meeting.title,
        meeting_date=meeting.meeting_date,
        speakers=list(names.items()),
        summary=protocol.summary,
        decisions=list(protocol.decisions),
        action_items=[{column.key: getattr(item, column.key) for column in item.__table__.columns} for item in items],
        transcript=[
            {
                "t": f"{int(segment.start_s) // 60:02d}:{int(segment.start_s) % 60:02d}",
                "speaker": names.get(segment.speaker_id, segment.speaker_id),
                "text": segment.text,
            }
            for segment in segments
        ],
    )
